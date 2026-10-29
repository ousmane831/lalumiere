from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import grades, models as m, reports, services as sv


def fmt(c):
    if isinstance(c, bool):
        return "Oui" if c else "Non"
    if isinstance(c, int):
        return f"{c:,}".replace(",", " ")
    return "" if c is None else str(c)


def to_xlsx(title, subtitle, sections):
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append([title]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle]); ws.append([])
    for sec in sections:
        if sec.get("title"):
            ws.append([sec["title"]]); ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.append(sec["headers"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F1C3F")
        for r in sec["rows"]:
            ws.append([("Oui" if x else "Non") if isinstance(x, bool) else x for x in r])
        ws.append([])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(40, max(12, *(len(str(c.value or "")) for c in col[3:]), 12) + 2)
    buf = BytesIO(); wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"


def to_pdf(title, subtitle, sections):
    school = m.SchoolSettings.get().name
    wide = max(len(s["headers"]) for s in sections) > 6
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4) if wide else A4, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    ss = getSampleStyleSheet()
    cell = ss["Normal"].clone("c", fontSize=8, leading=10)
    hcell = ss["Normal"].clone("h", fontSize=8, leading=10, textColor=colors.white, fontName="Helvetica-Bold")
    el = [Paragraph(f"<b>{escape(school)}</b> — {escape(title)}", ss["Title"]), Paragraph(escape(subtitle), ss["Normal"]), Spacer(1, 8)]
    for sec in sections:
        if sec.get("title"):
            el.append(Paragraph(f"<b>{escape(sec['title'])}</b>", ss["Heading3"]))
        data = [[Paragraph(escape(h), hcell) for h in sec["headers"]]] + [[Paragraph(escape(fmt(c)), cell) for c in r] for r in sec["rows"]]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f1c3f")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#bbc")),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f5fa")]), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        el += [t, Spacer(1, 10)]
    doc.build(el)
    return buf.getvalue(), "application/pdf", "pdf"


def sec(title, headers, rows):
    return {"title": title, "headers": headers, "rows": rows}


def build(kind, p):
    """Retourne (titre, sous-titre, sections) pour un type d'export."""
    month = sv.parse_month(p.get("month"))
    ml = month.strftime("%m/%Y")
    if kind == "students":
        qs = m.Student.objects.select_related("classroom")
        if p.get("classroom"):
            qs = qs.filter(classroom_id=p["classroom"])
        return "Liste des élèves", "", [sec("", ["Matricule", "Nom", "Prénom", "Classe", "Sexe", "Naissance", "Responsable", "Téléphone", "Mensualité", "Statut"],
            [[s.matricule, s.last_name, s.first_name, s.classroom.name, s.sex, s.birth_date and s.birth_date.strftime("%d/%m/%Y"), s.guardian_name, s.guardian_phone, s.monthly_fee, s.get_status_display()] for s in qs])]
    if kind == "dues":
        rows = sorted(sv.dues(month, p.get("classroom")), key=lambda r: -r["total_due"])
        lab = {"paid": "Payé", "partial": "Partiel", "none": "Impayé"}
        return f"Paiements et impayés — {ml}", "", [sec("", ["Élève", "Classe", "Téléphone", "Mensualité", "Payé", "Reste du mois", "Arriérés", "Total dû", "Statut"],
            [[r["name"], r["classroom"], r["guardian_phone"], r["monthly_fee"], r["paid"], r["remaining"], r["arrears"], r["total_due"], lab[r["status"]]] for r in rows])]
    if kind == "payments":
        qs = m.Payment.objects.filter(month=month).select_related("student__classroom", "received_by")
        return f"Paiements reçus — mensualité {ml}", "", [sec("", ["Reçu", "Élève", "Classe", "Mensualité", "Montant", "Date", "Mode", "Référence", "Caissier"],
            [[x.receipt_no, x.student.full_name, x.student.classroom.name, x.month.strftime("%m/%Y"), x.amount, x.paid_on.strftime("%d/%m/%Y"), x.get_method_display(), x.reference,
              (x.received_by.get_full_name() or x.received_by.username) if x.received_by else ""] for x in qs])]
    if kind in ("expenses", "incomes"):
        M = m.Expense if kind == "expenses" else m.Income
        qs = M.objects.filter(date__year=month.year, date__month=month.month)
        rows = [[x.date.strftime("%d/%m/%Y"), x.label, x.get_category_display(), getattr(x, "beneficiary", "") or getattr(x, "reference", ""), x.get_method_display(), x.amount] for x in qs]
        rows.append(["", "TOTAL", "", "", "", sum(x.amount for x in qs)])
        return ("Registre des dépenses" if kind == "expenses" else "Registre des recettes") + f" — {ml}", "", [sec("", ["Date", "Motif", "Catégorie", "Bénéficiaire / réf.", "Mode", "Montant"], rows)]
    if kind == "hours":
        qs = m.WorkedHours.objects.filter(month=month).select_related("teacher")
        from .serializers import WorkedHoursSerializer
        rows = []
        for w in qs:
            d = WorkedHoursSerializer(w).data
            rows.append([w.teacher.full_name, "Salaire fixe" if d["mode"] == "fixe" else f"{w.teacher.hourly_rate} FCFA/h", float(w.planned_hours), d["absent_hours"], d["replacement_hours"], float(w.done_hours), w.amount, w.paid])
        rows.append(["TOTAL", "", "", "", "", "", sum(w.amount for w in qs), ""])
        return f"État mensuel des heures et rémunérations — {ml}", "", [sec("", ["Enseignant", "Tarif", "Heures prévues", "Absences", "Remplacements", "Heures réalisées", "Montant", "Versé"], rows)]
    if kind == "results":
        room = m.ClassRoom.objects.get(pk=p["classroom"]); term = int(p.get("term", 1)); res = grades.class_results(room, term)
        st = m.SchoolSettings.get()
        rows = [[r["name"]] + [r["averages"].get(s["id"], "") for s in res["subjects"]] + ["" if r["general"] is None else r["general"], r["rank"] or ""] for r in res["rows"]]
        return f"Résultats {room.name} — {grades.term_name(st, term)}", f"Année {room.school_year}", [sec("", ["Élève"] + [s["name"] for s in res["subjects"]] + ["Moyenne", "Rang"], rows)]
    if kind == "report-monthly":
        r = reports.monthly_report(month)
        return f"Rapport mensuel — {ml}", "", [
            sec("Synthèse", ["Indicateur", "Montant (FCFA)"], [["Scolarité attendue", r["expected"]], ["Scolarité effectivement encaissée (sur les mensualités du mois)", r["collected"]],
                ["Impayés du mois", r["unpaid"]], ["Arriérés des mois précédents", r["arrears"]], ["Recettes du mois (toutes sources)", r["income"]], ["Dépenses payées", r["expenses"]],
                ["Salaires calculés", r["salaries_calculated"]], ["Salaires versés", r["salaries_paid"]], ["Solde de trésorerie", r["treasury"]]]),
            sec("Impayés par classe", ["Classe", "Élèves", "Attendu", "Encaissé", "Impayé du mois", "Arriérés"], [[c["classroom"], c["students"], c["expected"], c["collected"], c["unpaid"], c["arrears"]] for c in r["by_class"]]),
            sec("Recettes par mode de règlement", ["Mode", "Montant"], [[x["label"], x["total"]] for x in r["by_method"]]),
            sec("Dépenses par catégorie", ["Catégorie", "Montant"], [[x["label"], x["total"]] for x in r["by_category"]]),
            sec("Comptes (solde cumulé)", ["Compte", "Solde"], [[x["label"], x["balance"]] for x in r["accounts"]])]
    if kind == "report-annual":
        r = reports.annual_report(int(p.get("year") or month.year))
        rows = [[sv_label(x["month"]), x["expected"], x["collected"], x["unpaid"], x["income"], x["expenses"], x["balance"]] for x in r["months"]]
        t = r["totals"]; rows.append(["TOTAL", t["expected"], t["collected"], t["unpaid"], t["income"], t["expenses"], t["balance"]])
        return f"Rapport annuel {r['year']}", f"Trésorerie en fin d’année : {fmt(r['treasury'])} FCFA", [sec("", ["Mois", "Attendu", "Encaissé (scolarité)", "Impayés", "Recettes", "Dépenses", "Résultat"], rows)]
    raise KeyError(kind)


def sv_label(ym):
    names = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
    return f"{names[int(ym[5:]) - 1]} {ym[:4]}"


def render(kind, p, fmt_="xlsx"):
    title, sub, sections = build(kind, p)
    body, ctype, ext = (to_pdf if fmt_ == "pdf" else to_xlsx)(title, sub, sections)
    return body, ctype, f"{kind}.{ext}"
